"""
Bulk Content Generator (Elite Upgrade) — Optimized for Parallel AI Execution.
Reads top trending keywords and uses the AI Race to generate high-quality SEO blog posts.
"""

import asyncio
import json
import openpyxl
import os
from datetime import datetime
import pandas as pd

# Import our advanced model router
from core.model_router import race_models

class BulkContentGenerator:
    def __init__(self, tracker_path, output_dir="./seo_content"):
        self.tracker_path = tracker_path
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)
    
    async def query_ai(self, prompt, context="Content Gen"):
        """Execute a parallel AI race."""
        result = await race_models(prompt)
        return result.get("response", f"Error: No response from AI. {result.get('error', '')}")
    
    def extract_top_keywords(self, sheet_name="Trending Keywords", count=10):
        """Read top N keywords from the Excel tracker."""
        if not os.path.exists(self.tracker_path):
            print(f"⚠️ Tracker not found at {self.tracker_path}. Using sample keywords.")
            return [{"keyword": "painless dentistry hyderabad", "volume": 5000}]
            
        wb = openpyxl.load_workbook(self.tracker_path, data_only=True)
        if sheet_name not in wb.sheet_names:
            sheet_name = wb.sheet_names[0]
            
        ws = wb[sheet_name]
        keywords = []
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
            if not row or len(row) < 3: continue
            keyword = str(row[1]).strip() if row[1] else ""
            volume = str(row[2]).strip() if row[2] else "0"
            
            if keyword and keyword != "None":
                try:
                    vol = int(volume)
                except:
                    vol = 0
                keywords.append({"keyword": keyword, "volume": vol})
        
        keywords.sort(key=lambda x: x["volume"], reverse=True)
        return keywords[:count]
    
    async def generate_single_post(self, kw, i, total):
        """Generate a single blog post using the AI Race."""
        keyword = kw["keyword"]
        volume = kw["volume"]
        
        print(f"🦙 [{i+1}/{total}] Drafting: \"{keyword}\"...")
        
        prompt = f"""You are an expert SEO Content Writer for Cosmos Clinics, Hyderabad.
        Write a 500+ word SEO-optimized blog post for the keyword: "{keyword}" (Search Vol: {volume}).
        
        Requirements:
        1. Engaging H1, H2, H3 headers.
        2. Natural keyword integration.
        3. Local context (Hyderabad/Kukatpally).
        4. Meta Title & Meta Description.
        5. Markdown format.
        6. CTA for booking a smile makeover.
        """
        
        content = await self.query_ai(prompt, f"Post: {keyword}")
        
        # Save to file
        safe_name = keyword.lower().replace(" ", "_")[:50]
        filename = f"{safe_name}.md"
        filepath = os.path.join(self.output_dir, filename)
        
        with open(filepath, "w") as f:
            f.write(content)
            
        return {"keyword": keyword, "file": filepath, "length": len(content)}

    async def run_parallel(self, count=5):
        """Generate multiple posts in parallel leveraging the AI Arsenal."""
        print(f"📝 Bulk Content Generator (Elite Upgrade) — AI Parallel Mode")
        print(f"📂 Tracker: {self.tracker_path}\n")
        
        keywords = self.extract_top_keywords(count=count)
        if not keywords:
            print("❌ No keywords found. Aborting.")
            return []
            
        print(f"🎯 Preparing {len(keywords)} posts simultaneously...")
        tasks = [self.generate_single_post(kw, i, len(keywords)) for i, kw in enumerate(keywords)]
        
        results = await asyncio.gather(*tasks)
        
        print(f"\n{'='*50}")
        print(f"📊 PARALLEL CONTENT GENERATION COMPLETE")
        print(f"{'='*50}")
        for res in results:
            print(f"  ✅ {res['keyword']} -> {res['file']} ({res['length']} chars)")
            
        return results

if __name__ == "__main__":
    tracker = "./cosmos_marketing_tracker.xlsx"
    generator = BulkContentGenerator(tracker)
    asyncio.run(generator.run_parallel(count=3))
